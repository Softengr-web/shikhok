#!/usr/bin/env python3
"""Audit-first DOC/DOCX/PDF question importer for Private Tutor.

Text is kept source-faithful. This tool never rewrites question wording. Every
detected numbered candidate is sent as either an import candidate or a review
item; unreadable files and pages are recorded in the file audit.
"""
from __future__ import annotations

import argparse
import base64
import csv
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from queue import Queue
from threading import Thread, local
import time
import urllib.error
import urllib.request
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable
from zipfile import BadZipFile, ZipFile

SUPPORTED = {".doc", ".docx", ".pdf"}
BN_DIGITS = str.maketrans("০১২৩৪৫৬৭৮৯", "0123456789")
BN_LABELS = {"ক": 0, "খ": 1, "গ": 2, "ঘ": 3}
EN_LABELS = {"a": 0, "b": 1, "c": 2, "d": 3}
OPTION_MARKER = re.compile(r"(?<!\S)(?:\((?P<paren>[কখগঘa-dA-D])\)|(?P<plain>[কখগঘa-dA-D])[).।:])(?=\s|$)")
QUESTION_START = re.compile(
    r"^\s*(?:(?:প্রশ্ন|question|q)\s*)?(?:\(?\s*)"
    r"(?P<number>[0-9০-৯]{1,4})\s*(?:[.)।:ঃ]|\))\s*(?P<text>.*)$",
    re.IGNORECASE,
)
ANSWER = re.compile(
    r"(?:সঠিক\s*)?(?:উত্তর|answer|ans)\s*[:：ঃ=\-]?\s*\(?\s*([কখগঘa-dA-D১-৪1-4])\s*\)?",
    re.IGNORECASE,
)
KEY_HEADING = re.compile(r"(?:উত্তরমালা|উত্তর\s*সূচি|answer\s*key|answer\s*sheet)", re.IGNORECASE)
KEY_PAIR = re.compile(r"(?<!\S)([0-9০-৯]{1,4})\s*[.)।:ঃ-]\s*\(?\s*([কখগঘa-dA-D])\s*\)?")
ANSWER_BULLET = "●"
BIJOY_FONT = re.compile(r"(?:sutonny|bijoy|sulekha|boishakh|kongsho|(?:[a-z0-9_-]+mj))", re.IGNORECASE)
WORD_LOCAL = local()


@dataclass
class PageText:
    page: int | None
    text: str
    confidence: float | None = None
    media: list[dict[str, Any]] | None = None
    error: str | None = None


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_number(value: str) -> str:
    return value.translate(BN_DIGITS)


def option_index(label: str) -> int | None:
    value = label.strip("() ).।:").casefold()
    if value in BN_LABELS:
        return BN_LABELS[value]
    if value in EN_LABELS:
        return EN_LABELS[value]
    if value.isdigit():
        return int(normalize_number(value)) - 1
    return None


def split_inline_options(line: str) -> tuple[str, list[tuple[str, str]]] | None:
    matches = list(OPTION_MARKER.finditer(line))
    if len(matches) < 2:
        return None
    labels = [(match.group("paren") or match.group("plain") or "").casefold() for match in matches]
    if len(set(labels)) != len(labels):
        return None
    prefix = line[: matches[0].start()].rstrip()
    options: list[tuple[str, str]] = []
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(line)
        label = match.group("paren") or match.group("plain") or ""
        options.append((label, line[match.end() : end].strip()))
    return prefix, options


def _candidate_ranges(lines: list[str]) -> tuple[list[tuple[str, list[str], int, str]], dict[str, str], int | None]:
    key_start = next((index for index, line in enumerate(lines) if KEY_HEADING.search(line)), None)
    content_end = key_start if key_start is not None else len(lines)
    answer_key: dict[str, str] = {}
    if key_start is not None:
        collected: dict[str, set[str]] = {}
        for line in lines[key_start:]:
            for match in KEY_PAIR.finditer(line):
                collected.setdefault(normalize_number(match.group(1)), set()).add(match.group(2))
        answer_key = {number: next(iter(labels)) for number, labels in collected.items() if len(labels) == 1}

    starts: list[tuple[int, str, str]] = []
    for index, line in enumerate(lines[:content_end]):
        match = QUESTION_START.match(line)
        if match:
            starts.append((index, normalize_number(match.group("number")), match.group("text")))
    ranges: list[tuple[str, list[str], int, str]] = []
    for position, (start, number, first_text) in enumerate(starts):
        end = starts[position + 1][0] if position + 1 < len(starts) else content_end
        candidate_lines = [first_text, *lines[start + 1 : end]]
        ranges.append((number, candidate_lines, start + 1, "numbered"))
    if starts:
        return ranges, answer_key, key_start

    # Some study sheets omit question numbers. Recover paragraphs which contain
    # at least two recognized option labels, keeping each block as a candidate.
    blocks: list[list[str]] = []
    block: list[str] = []
    for line in lines[:content_end]:
        if not line.strip():
            if block:
                blocks.append(block)
                block = []
        else:
            block.append(line)
    if block:
        blocks.append(block)
    for index, candidate_lines in enumerate(blocks, start=1):
        if any(split_inline_options(line) for line in candidate_lines):
            ranges.append((str(index), candidate_lines, 1, "un-numbered"))
    return ranges, answer_key, key_start


def parse_mcqs(page: PageText, shared_answer_key: dict[str, str] | None = None, shared_answer_key_ocr: bool = False) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    lines = page.text.splitlines()
    ranges, page_answer_key, _ = _candidate_ranges(lines)
    answer_key = {**(shared_answer_key or {}), **page_answer_key}
    questions: list[dict[str, Any]] = []
    issues: list[dict[str, Any]] = []
    for number, candidate_lines, source_line, numbering in ranges:
        shared_answer_used = number in (shared_answer_key or {}) and number not in page_answer_key
        question_lines: list[str] = []
        option_parts: list[tuple[str, str]] = []
        answer_label: str | None = None
        for line in candidate_lines:
            if ANSWER_BULLET in line:
                before, after = line.split(ANSWER_BULLET, 1)
                before = before.strip()
                after = after.strip()
                if before:
                    inline = split_inline_options(before)
                    if inline:
                        prefix, found_options = inline
                        if prefix:
                            question_lines.append(prefix)
                        option_parts.extend(found_options)
                    else:
                        marker = OPTION_MARKER.match(before)
                        if marker:
                            option_parts.append((marker.group("paren") or marker.group("plain") or "", before[marker.end():].strip()))
                        elif option_parts:
                            label, text = option_parts[-1]
                            option_parts[-1] = (label, f"{text}\n{before}".strip())
                        else:
                            question_lines.append(before)
                following_options = list(OPTION_MARKER.finditer(after))
                if following_options and following_options[0].start() == 0:
                    answer_option = following_options[0]
                    answer_label = answer_option.group("paren") or answer_option.group("plain") or ""
                    answer_end = following_options[1].start() if len(following_options) > 1 else len(after)
                    option_parts.append((answer_label, after[answer_option.end():answer_end].strip()))
                    following_options = following_options[1:]
                else:
                    used = {option_index(label) for label, _ in option_parts}
                    inferred_index = next((index for index in range(4) if index not in used), None)
                    if inferred_index is not None:
                        answer_label = "কখগঘ"[inferred_index]
                        answer_end = following_options[0].start() if following_options else len(after)
                        option_parts.append((answer_label, after[:answer_end].strip()))
                    else:
                        issues.append({"source_page": page.page, "source_line": source_line, "raw_text": line, "problem": "উত্তর চিহ্নিত ● পাওয়া গেছে কিন্তু option label নির্ভরযোগ্যভাবে মেলেনি", "suggested_correction": "মূল ফাইল দেখে সঠিক option চিহ্নিত করুন।", "interpretation": {"answer_marker": ANSWER_BULLET}})
                for option_index_in_line, marker in enumerate(following_options):
                    end = following_options[option_index_in_line + 1].start() if option_index_in_line + 1 < len(following_options) else len(after)
                    label = marker.group("paren") or marker.group("plain") or ""
                    option_parts.append((label, after[marker.end():end].strip()))
                continue
            answer_match = ANSWER.search(line)
            if answer_match:
                answer_label = answer_match.group(1)
                remainder = ANSWER.sub("", line).strip(" \t:ঃ-–—")
                if remainder:
                    question_lines.append(remainder)
                continue
            split = split_inline_options(line)
            if split:
                prefix, inline = split
                if prefix:
                    (option_parts[-1] if option_parts else None)
                    if option_parts:
                        last_label, last_text = option_parts[-1]
                        option_parts[-1] = (last_label, f"{last_text}\n{prefix}".strip())
                    else:
                        question_lines.append(prefix)
                option_parts.extend(inline)
                continue
            marker = OPTION_MARKER.match(line)
            if marker:
                label = marker.group("paren") or marker.group("plain") or ""
                option_parts.append((label, line[marker.end() :].strip()))
            elif option_parts:
                label, text = option_parts[-1]
                option_parts[-1] = (label, f"{text}\n{line.strip()}".strip())
            else:
                question_lines.append(line)

        if answer_label is None:
            answer_label = answer_key.get(number)
        question_text = "\n".join(part for part in question_lines if part.strip()).strip()
        options = [{"label": label, "text": text.strip(), "media_sha256": re.findall(r"\[ছবি\s+([a-f0-9]{64})\]", text, re.IGNORECASE)} for label, text in option_parts]
        if options and all(option_index(option["label"]) is not None for option in options):
            options.sort(key=lambda option: option_index(option["label"]) or 0)
        answer = option_index(answer_label) if answer_label is not None else None
        problems: list[str] = []
        if not question_text:
            problems.append("প্রশ্নের মূল বক্তব্য শনাক্ত হয়নি")
        if len(options) < 2:
            problems.append(f"কমপক্ষে ২টি option শনাক্ত হয়নি (পাওয়া গেছে {len(options)}টি)")
        if any(not option["text"] for option in options):
            problems.append("এক বা একাধিক option-এর লেখা ফাঁকা")
        if answer is None:
            problems.append("সঠিক উত্তরের key পাওয়া যায়নি")
        elif answer < 0 or answer >= len(options):
            problems.append(f"উত্তর key {answer_label!r} মেলা option নয়")
        if page.confidence is not None:
            if page.confidence < 70:
                problems.append(f"OCR confidence কম ({page.confidence:.0f}%)")
            else:
                problems.append(f"OCR text মূল পৃষ্ঠার সঙ্গে হাতে যাচাই করতে হবে (confidence {page.confidence:.0f}%)")
        if shared_answer_key_ocr and shared_answer_used:
            problems.append("separate answer key OCR থেকে এসেছে; মূল উত্তরমালা দেখে হাতে যাচাই করুন")
        raw_candidate = "\n".join(candidate_lines).strip()
        media_hashes = set(re.findall(r"\[ছবি\s+([a-f0-9]{64})\]", raw_candidate, re.IGNORECASE))
        if page.page is not None and page.media:
            media_hashes.update(item["sha256"] for item in page.media)
        common = {
            "source_number": number,
            "source_page": page.page,
            "source_line": source_line,
            "source_location": f"লাইন {source_line}" if page.page is None else f"পৃষ্ঠা {page.page}, লাইন {source_line}",
            "numbering_style": numbering,
            "raw_text": raw_candidate,
            "media_sha256": sorted(media_hashes),
        }
        question_text = re.sub(r"\[ছবি\s+[a-f0-9]{64}\]", "", question_text, flags=re.IGNORECASE).strip()
        options = [{**option, "text": re.sub(r"\[ছবি\s+[a-f0-9]{64}\]", "", option["text"], flags=re.IGNORECASE).strip()} for option in options]
        if problems:
            issues.append({**common, "problem": "; ".join(problems), "suggested_correction": "মূল ফাইল দেখে প্রশ্ন, option ও answer যাচাই করুন।", "interpretation": {"question": question_text, "options": options, "answer_label": answer_label}})
        else:
            questions.append({**common, "question": question_text, "options": options, "correct_option": answer, "explanation": ""})
    if page.error:
        issues.append({"source_page": page.page, "raw_text": page.text, "problem": page.error, "suggested_correction": "পৃষ্ঠা/ফাইলটি হাতে পরীক্ষা করে আবার import করুন।", "interpretation": {"extraction_error": page.error}})
    if not ranges and (OPTION_MARKER.search(page.text) or page.text.strip()):
        if OPTION_MARKER.search(page.text):
            issues.append({"source_page": page.page, "raw_text": page.text[:4000], "problem": "option label পাওয়া গেছে কিন্তু question boundary নির্ভরযোগ্যভাবে শনাক্ত হয়নি", "suggested_correction": "প্রশ্নগুলো আলাদা করে review queue-তে চিহ্নিত করুন।", "interpretation": {"option_markers_found": True}})
    return questions, issues


def _document_answer_key(pages: list[PageText]) -> tuple[dict[str, str], bool]:
    key_sections: list[tuple[dict[str, str], bool]] = []
    for page in pages:
        lines = page.text.splitlines()
        if any(KEY_HEADING.search(line) for line in lines):
            _, answer_key, _ = _candidate_ranges(lines)
            key_sections.append((answer_key, page.confidence is not None))
    if len(key_sections) != 1:
        return {}, False
    return key_sections[0]


def _render_page_image(page: Any, max_width: int = 1600) -> bytes:
    import pymupdf as fitz  # type: ignore
    from PIL import Image  # type: ignore
    from io import BytesIO

    pixmap = page.get_pixmap(matrix=fitz.Matrix(1.6, 1.6), alpha=False)
    image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
    if image.width > max_width:
        height = max(1, round(image.height * max_width / image.width))
        image = image.resize((max_width, height), Image.Resampling.LANCZOS)
    output = BytesIO()
    image.save(output, format="JPEG", quality=78, optimize=True)
    return output.getvalue()


def _convert_bijoy_texts(texts: list[str]) -> list[str]:
    if not texts:
        return []
    node = shutil.which("node")
    converter = Path(__file__).with_name("bijoy-convert.mjs")
    if not node or not converter.exists():
        raise RuntimeError("Legacy Bijoy/Sutonny font found; Node.js and the bijoy2unicode helper are required to convert it safely")
    result = subprocess.run(
        [node, str(converter), "--texts"], input=json.dumps(texts, ensure_ascii=False),
        capture_output=True, text=True, encoding="utf-8", timeout=120
    )
    if result.returncode != 0:
        raise RuntimeError(f"Bijoy text conversion failed: {(result.stderr or result.stdout).strip()[:800]}")
    converted = json.loads(result.stdout)
    if len(converted) != len(texts) or any(not isinstance(text, str) for text in converted):
        raise RuntimeError("Bijoy conversion returned an incomplete text list")
    return converted


def _optimize_image(data: bytes, max_width: int = 1600) -> tuple[bytes, str]:
    try:
        from PIL import Image
        from io import BytesIO

        image = Image.open(BytesIO(data))
        image.load()
        if image.width > max_width:
            height = max(1, round(image.height * max_width / image.width))
            image = image.resize((max_width, height), Image.Resampling.LANCZOS)
        if image.mode not in ("RGB", "L"):
            background = Image.new("RGB", image.size, "white")
            if "A" in image.getbands():
                background.paste(image.convert("RGBA"), mask=image.convert("RGBA").getchannel("A"))
            else:
                background.paste(image.convert("RGB"))
            image = background
        output = BytesIO()
        if len(data) < 1_000_000 and image.format in ("PNG", "JPEG"):
            image.save(output, format=image.format, optimize=True)
            mime_type = "image/png" if image.format == "PNG" else "image/jpeg"
        else:
            image.save(output, format="JPEG", quality=85, optimize=True)
            mime_type = "image/jpeg"
        encoded = output.getvalue()
        if len(encoded) <= 2_000_000:
            return encoded, mime_type
    except Exception:
        pass
    if len(data) <= 2_000_000:
        return data, "application/octet-stream"
    raise RuntimeError("image is larger than 2 MB after compression; question has been sent to review")


def _ocr_pdf_page(page: Any, languages: str) -> tuple[str, float | None, str | None]:
    try:
        import pymupdf as fitz  # type: ignore
        import pytesseract  # type: ignore
        from PIL import Image
        from io import BytesIO

        if not shutil.which("tesseract") and os.environ.get("LOCALAPPDATA"):
            user_install = Path(os.environ["LOCALAPPDATA"]) / "Programs" / "Tesseract-OCR" / "tesseract.exe"
            if user_install.is_file():
                pytesseract.pytesseract.tesseract_cmd = str(user_install)

        pixmap = page.get_pixmap(matrix=fitz.Matrix(2.0, 2.0), alpha=False)
        image = Image.open(BytesIO(pixmap.tobytes("png")))
        data = pytesseract.image_to_data(image, lang=languages, output_type=pytesseract.Output.DICT, config="--psm 6")
        chunks: list[str] = []
        confidences: list[float] = []
        last_line: tuple[int, int, int] | None = None
        for index, token in enumerate(data.get("text", [])):
            value = str(token).strip()
            if not value:
                continue
            conf = float(data["conf"][index])
            if conf >= 0:
                confidences.append(conf)
            line_key = (int(data["block_num"][index]), int(data["par_num"][index]), int(data["line_num"][index]))
            if last_line is not None and last_line != line_key:
                chunks.append("\n")
            elif chunks and chunks[-1] != "\n":
                chunks.append(" ")
            chunks.append(value)
            last_line = line_key
        text = "".join(chunks).strip()
        confidence = sum(confidences) / len(confidences) if confidences else None
        if not text:
            return "", confidence, f"পৃষ্ঠা {page.number + 1}-এ text layer নেই এবং বাংলা OCR কোনো লেখা ফেরত দেয়নি"
        return text, confidence, None
    except Exception as error:  # missing Tesseract or Bengali trained data is auditable
        return "", None, f"পৃষ্ঠা {page.number + 1} scanned/image-only; OCR করা যায়নি: {error}"


def extract_pdf(path: Path, ocr_languages: str) -> list[PageText]:
    import pymupdf as fitz  # type: ignore

    document = fitz.open(path)
    extracted: list[dict[str, Any]] = []
    bijoy_texts: list[str] = []
    for index, page in enumerate(document):
        page_lines: list[dict[str, Any]] = []
        for block in page.get_text("dict", sort=False).get("blocks", []):
            for line in block.get("lines", []):
                spans: list[dict[str, Any]] = []
                for span in line.get("spans", []):
                    span_text = str(span.get("text", ""))
                    is_bijoy = bool(BIJOY_FONT.search(str(span.get("font", ""))))
                    span_index = len(bijoy_texts) if is_bijoy and span_text else None
                    if span_index is not None:
                        bijoy_texts.append(span_text)
                    spans.append({"text": span_text, "bijoy_index": span_index})
                if spans:
                    page_lines.append({"x": min(float(span["bbox"][0]) for span in line["spans"]), "y": float(line["bbox"][1]), "spans": spans})
        extracted.append({"page": page, "number": index + 1, "lines": page_lines})

    converted_bijoy = _convert_bijoy_texts(bijoy_texts)
    pages: list[PageText] = []
    for item in extracted:
        page = item["page"]
        text_columns: list[list[tuple[float, float, str]]] = [[], []]
        for line in item["lines"]:
            parts = [converted_bijoy[span["bijoy_index"]] if span["bijoy_index"] is not None else span["text"] for span in line["spans"]]
            line_text = "".join(parts).strip()
            if line_text:
                column = 0 if line["x"] < page.rect.width / 2 else 1
                text_columns[column].append((line["y"], line["x"], line_text))
        ordered_lines: list[str] = []
        for column in text_columns:
            merged: list[dict[str, Any]] = []
            for y, x, line_text in sorted(column, key=lambda row: (row[0], row[1])):
                if merged and abs(y - merged[-1]["y"]) <= 2.5:
                    merged[-1]["items"].append((x, line_text))
                else:
                    merged.append({"y": y, "items": [(x, line_text)]})
            ordered_lines.extend(" ".join(text for _, text in sorted(row["items"], key=lambda part: part[0])) for row in merged)
        text = "\n".join(ordered_lines).strip()
        confidence: float | None = None
        error: str | None = None
        media: list[dict[str, Any]] | None = None
        if not text:
            text, confidence, error = _ocr_pdf_page(page, ocr_languages)
        if page.get_images(full=True) or error or confidence is not None:
            try:
                data = _render_page_image(page)
                media = [{"page": item["number"], "mime_type": "image/jpeg", "sha256": sha256_bytes(data), "base64": base64.b64encode(data).decode("ascii")}]
            except Exception as image_error:
                error = (error + "; " if error else "") + f"পৃষ্ঠা ছবি সংরক্ষণ করা যায়নি: {image_error}"
        pages.append(PageText(item["number"], text, confidence, media, error))
    document.close()
    return pages


def extract_docx(path: Path) -> list[PageText]:
    from docx import Document  # type: ignore

    if _docx_uses_bijoy_font(path):
        node = shutil.which("node")
        converter = Path(__file__).with_name("bijoy-convert.mjs")
        if not node or not converter.exists():
            raise RuntimeError("Legacy Bijoy/Sutonny font found; Node.js and the bijoy2unicode helper are required to convert it safely")
        with tempfile.TemporaryDirectory(prefix="mcq-bijoy-") as tmp:
            converted = Path(tmp) / path.name
            result = subprocess.run([node, str(converter), "--docx", str(path), str(converted)], capture_output=True, text=True, timeout=180)
            if result.returncode != 0 or not converted.exists():
                raise RuntimeError(f"Bijoy DOCX conversion failed: {(result.stderr or result.stdout).strip()[:800]}")
            return _extract_docx_content(Document(converted))
    return _extract_docx_content(Document(path))


def _docx_uses_bijoy_font(path: Path) -> bool:
    try:
        with ZipFile(path) as archive:
            for name in archive.namelist():
                if name.endswith(".xml") and BIJOY_FONT.search(archive.read(name).decode("utf-8", errors="ignore")):
                    return True
    except (BadZipFile, OSError):
        return False
    return False


def _extract_docx_content(document: Any) -> list[PageText]:
    from docx.table import Table  # type: ignore
    from docx.text.paragraph import Paragraph  # type: ignore

    image_map: dict[str, dict[str, Any]] = {}
    for rel in document.part.rels.values():
        if rel.reltype.endswith("/image"):
            blob, content_type = _optimize_image(rel.target_part.blob)
            image_map[rel.rId] = {"page": 0, "mime_type": content_type, "sha256": sha256_bytes(blob), "base64": base64.b64encode(blob).decode("ascii")}

    def paragraph_text(paragraph: Any) -> str:
        text = paragraph.text
        relation_ids = [*paragraph._p.xpath(".//a:blip/@r:embed"), *paragraph._p.xpath(".//*[local-name()='imagedata']/@*[local-name()='id']")]
        markers = [f"[ছবি {image_map[rid]['sha256']}]" for rid in dict.fromkeys(relation_ids) if rid in image_map]
        return f"{text} {' '.join(markers)}".strip() if markers else text

    parts: list[str] = []
    location = 0
    for item in document.iter_inner_content():
        location += 1
        if isinstance(item, Paragraph):
            parts.append(paragraph_text(item))
        elif isinstance(item, Table):
            for row in item.rows:
                parts.append("\t".join(" / ".join(paragraph_text(paragraph) for paragraph in cell.paragraphs) for cell in row.cells))
    # First-order image references are kept in the source text; the import API
    # stores each asset once and links only assets referenced by a candidate.
    return [PageText(None, "\n".join(parts), media=list(image_map.values()))]


def extract_doc(path: Path) -> list[PageText]:
    office = shutil.which("soffice") or shutil.which("libreoffice")
    antiword = shutil.which("antiword")
    if office:
        with tempfile.TemporaryDirectory(prefix="mcq-doc-") as tmp:
            result = subprocess.run([office, "--headless", "--convert-to", "docx", "--outdir", tmp, str(path)], capture_output=True, text=True, timeout=180)
            converted = Path(tmp) / f"{path.stem}.docx"
            if result.returncode == 0 and converted.exists():
                return extract_docx(converted)
            detail = (result.stderr or result.stdout).strip()
            raise RuntimeError(f"LibreOffice .doc conversion failed: {detail or result.returncode}")
    if antiword:
        result = subprocess.run([antiword, str(path)], capture_output=True, timeout=120)
        if result.returncode == 0:
            return [PageText(None, result.stdout.decode("utf-8", errors="replace"))]
        raise RuntimeError(f"antiword conversion failed: {result.stderr.decode('utf-8', errors='replace')[:500]}")
    try:
        import win32com.client  # type: ignore

        last_error: Exception | None = None
        for attempt in range(2):
            try:
                if getattr(WORD_LOCAL, "app", None) is None:
                    WORD_LOCAL.app = win32com.client.DispatchEx("Word.Application")
                    WORD_LOCAL.app.Visible = False
                    WORD_LOCAL.app.DisplayAlerts = 0
                    WORD_LOCAL.app.AutomationSecurity = 3  # Force-disable macros for source documents.
                with tempfile.TemporaryDirectory(prefix="mcq-word-") as tmp:
                    converted = Path(tmp) / f"{path.stem}.docx"
                    document = WORD_LOCAL.app.Documents.Open(str(path), ReadOnly=True, ConfirmConversions=False, AddToRecentFiles=False, Visible=False)
                    try:
                        document.SaveAs2(str(converted), FileFormat=16)  # wdFormatDocumentDefault (.docx)
                    finally:
                        document.Close(SaveChanges=0)
                    return extract_docx(converted)
            except Exception as error:
                last_error = error
                # Word can exit after a corrupt legacy document. A stale COM
                # proxy then fails every later .doc, so restart it once.
                close_word()
                if attempt == 1:
                    raise
        raise RuntimeError("Word COM conversion failed") from last_error
    except Exception as error:
        raise RuntimeError(f".doc could not be converted; install LibreOffice/antiword or enable Microsoft Word COM: {error}") from error


def close_word() -> None:
    app = getattr(WORD_LOCAL, "app", None)
    if app is not None:
        try:
            app.Quit(SaveChanges=0)
        except Exception:
            pass
        WORD_LOCAL.app = None


def extract_file(path: Path, ocr_languages: str) -> list[PageText]:
    suffix = path.suffix.casefold()
    if suffix == ".pdf":
        return extract_pdf(path, ocr_languages)
    if suffix == ".docx":
        return extract_docx(path)
    if suffix == ".doc":
        return extract_doc(path)
    raise ValueError(f"unsupported file extension: {suffix}")


def _attach_pdf_issue_previews(path: Path, issues: list[dict[str, Any]], media_by_hash: dict[str, dict[str, Any]]) -> None:
    page_numbers = {int(issue["source_page"]) for issue in issues if issue.get("source_page") and str(issue["source_page"]).isdigit()}
    if not page_numbers:
        return
    import pymupdf  # type: ignore

    document = pymupdf.open(path)
    try:
        for page_number in sorted(page_numbers):
            if any(int(media.get("page") or 0) == page_number for media in media_by_hash.values()):
                continue
            if page_number < 1 or page_number > len(document):
                continue
            data = _render_page_image(document[page_number - 1])
            digest = sha256_bytes(data)
            media_by_hash[digest] = {"page": page_number, "mime_type": "image/jpeg", "sha256": digest, "base64": base64.b64encode(data).decode("ascii")}
            for issue in issues:
                if issue.get("source_page") == page_number:
                    issue["media_sha256"] = sorted(set(issue.get("media_sha256", [])) | {digest})
    finally:
        document.close()


def _clean_folder(value: str) -> str:
    return re.sub(r"^\s*[0-9০-৯]+[.\s-]+", "", value).strip(" #")


def infer_metadata(path: Path, root: Path) -> dict[str, str | None]:
    relative = path.relative_to(root)
    folders = list(relative.parts[:-1])
    class_level = "Unknown"
    group: str | None = None
    class_root_index = 0
    match = re.search(r"class\s*(\d+)(?:\s*[-–]\s*(\d+))?", root.name, re.IGNORECASE)
    class_folder_is_root = bool(match)
    root_name = root.name if match else (folders[0] if folders else "")
    if not match:
        match = re.search(r"class\s*(\d+)(?:\s*[-–]\s*(\d+))?", root_name, re.IGNORECASE)
    if match:
        first, last = match.group(1), match.group(2)
        class_level = f"Class {first}-{last}" if last else f"Class {first}"
        class_root_index = 0 if class_folder_is_root else 1
        part_match = re.search(r"part\s*(\d+)", root_name, re.IGNORECASE)
        if part_match:
            group = "SSC"
    elif re.search(r"HSC", root.name, re.IGNORECASE) or re.search(r"HSC", root_name, re.IGNORECASE):
        class_level = "Class 11-12"
        class_root_index = 0 if re.search(r"HSC", root.name, re.IGNORECASE) else 1
        group = "HSC"

    descendants = folders[class_root_index:]
    if descendants:
        first = _clean_folder(descendants[0])
        if re.search(r"(?:SSC|HSC)\s+(?:common|science|humanit|business|commerce)|common\s+subject", first, re.IGNORECASE):
            group = first
            subject_folder = descendants[1] if len(descendants) > 1 else first
        else:
            subject_folder = descendants[0]
        subject = _clean_folder(subject_folder)
    else:
        subject = "Unknown"

    part: str | None = None
    for folder in folders:
        part_match = re.search(r"part\s*(\d+)", folder, re.IGNORECASE)
        if part_match:
            part = f"Part {part_match.group(1)}"
        paper = re.search(r"(\d+)(st|nd|rd|th)\s+paper", folder, re.IGNORECASE)
        if paper:
            part = f"{paper.group(1)}{paper.group(2).casefold()} Paper"
    if part is None and any(re.search(r"\bmcq\b", folder, re.IGNORECASE) for folder in folders):
        part = "MCQ"
    chapter: str | None = None
    chapter_source = " ".join([*folders, path.stem])
    chapter_match = re.search(r"(?:chapter|chpater|ch\.?|অধ্যায়|অধ্যায়|পাঠ)\s*[-_ ]*([0-9০-৯]+)", chapter_source, re.IGNORECASE)
    if chapter_match:
        chapter = f"Chapter {int(normalize_number(chapter_match.group(1)))}"
    elif re.search(r"chapterwise|অধ্যায়ভিত্তিক", chapter_source, re.IGNORECASE):
        chapter = _clean_folder(path.stem)
    elif any(re.search(r"\bmcq\b", folder, re.IGNORECASE) for folder in folders):
        # Many lesson-specific packs name the chapter in the filename rather
        # than in a directory (for example, `2. MCQ/4. Tolpar.pdf`).
        chapter = _clean_folder(path.stem)
    return {"class_level": class_level, "group_name": group, "subject": subject, "part": part, "chapter": chapter, "topic": None}


def process_file(path: Path, root: Path, ocr_languages: str) -> dict[str, Any]:
    digest = sha256_file(path)
    metadata = infer_metadata(path, root)
    source = {
        "source_path": path.relative_to(root).as_posix(),
        "source_filename": path.name,
        "file_type": path.suffix[1:].upper(),
        "file_size": path.stat().st_size,
        "file_hash": digest,
        **metadata,
    }
    pages = extract_file(path, ocr_languages)
    shared_answer_key, shared_answer_key_ocr = _document_answer_key(pages)
    questions: list[dict[str, Any]] = []
    issues: list[dict[str, Any]] = []
    media_by_hash: dict[str, dict[str, Any]] = {}
    next_candidate = 0
    for page in pages:
        if page.media:
            for media in page.media:
                media_by_hash[media["sha256"]] = media
        found, review = parse_mcqs(page, shared_answer_key, shared_answer_key_ocr)
        for candidate in found:
            next_candidate += 1
            questions.append({**candidate, "candidate_index": next_candidate})
        for issue in review:
            next_candidate += 1
            issues.append({**issue, "candidate_index": next_candidate})
    if path.suffix.casefold() == ".pdf":
        _attach_pdf_issue_previews(path, issues, media_by_hash)
    if not pages:
        issues.append({"candidate_index": None, "source_page": None, "raw_text": "", "problem": "ফাইলে কোনো page/text পাওয়া যায়নি", "suggested_correction": "মূল ফাইল খুলে পরীক্ষা করুন।", "interpretation": {}})
    elif not questions and not issues:
        issues.append({"candidate_index": None, "source_page": None, "raw_text": "\n".join(page.text for page in pages)[:4000], "problem": "কোনো MCQ candidate শনাক্ত হয়নি; ফাইলটি নীরবে বাদ না দিয়ে review-তে রাখা হয়েছে", "suggested_correction": "ফাইলের layout ও question format পরীক্ষা করুন।", "interpretation": {"pages": len(pages)}})
    return {"source": source, "questions": questions, "issues": issues, "media": list(media_by_hash.values()), "page_count": len(pages)}


def _import_worker(tasks: Queue, results: Queue, root: Path, ocr_languages: str) -> None:
    """Keep one Word COM server per worker; isolate it from other apartments."""
    com_runtime = None
    try:
        try:
            import pythoncom  # type: ignore

            com_runtime = pythoncom
            com_runtime.CoInitialize()
        except ImportError:
            pass
        while True:
            task = tasks.get()
            if task is None:
                return
            index, path = task
            started = time.monotonic()
            try:
                results.put((index, process_file(path, root, ocr_languages), None, time.monotonic() - started))
            except Exception as error:
                results.put((index, None, error, time.monotonic() - started))
    finally:
        close_word()
        if com_runtime is not None:
            com_runtime.CoUninitialize()


class ApiClient:
    def __init__(self, base_url: str, email: str, password: str):
        self.base_url = base_url.rstrip("/")
        self.cookie = ""
        self.email = email
        self.password = password
        payload = self.request("POST", "/auth/login", {"email": email, "password": password}, login=True)
        if payload.get("role") not in ("ADMIN", "SUPER_ADMIN"):
            raise RuntimeError("Import API login account must have ADMIN or SUPER_ADMIN role")

    def request(self, method: str, path: str, body: Any, login: bool = False) -> Any:
        data = json.dumps(body, ensure_ascii=False).encode("utf-8")
        request = urllib.request.Request(self.base_url + "/api" + path, data=data, method=method, headers={"Content-Type": "application/json", **({"Cookie": self.cookie} if self.cookie else {})})
        try:
            with urllib.request.urlopen(request, timeout=90) as response:
                if login:
                    self.cookie = response.headers.get("Set-Cookie", "").split(";", 1)[0]
                payload = json.loads(response.read().decode("utf-8"))
        except urllib.error.HTTPError as error:
            detail = error.read().decode("utf-8", errors="replace")[:1000]
            if error.code == 401 and not login:
                self.cookie = ""
                self.request("POST", "/auth/login", {"email": self.email, "password": self.password}, login=True)
                return self.request(method, path, body)
            raise RuntimeError(f"API {method} {path} returned {error.code}: {detail}") from error
        if not payload.get("ok"):
            raise RuntimeError(payload.get("message", "API request failed"))
        return payload.get("data")

    def send_file(self, record: dict[str, Any], batch_size: int = 20) -> dict[str, Any]:
        candidates = [{"candidate_index": item["candidate_index"], "question": item} for item in record["questions"]]
        candidates += [{"candidate_index": item.get("candidate_index"), "issue": item} for item in record["issues"] if item.get("candidate_index") is not None]
        if not candidates:
            candidates = []
        media_by_hash = {item["sha256"]: item for item in record["media"]}
        for offset in range(0, max(1, len(candidates)), batch_size):
            batch = candidates[offset : offset + batch_size]
            used_hashes = {digest for item in batch for digest in (item.get("question") or item.get("issue") or {}).get("media_sha256", [])}
            if offset == 0 and not batch:
                used_hashes = set(media_by_hash)
            assets = [media_by_hash[key] for key in used_hashes if key in media_by_hash]
            source_issues = [item for item in record["issues"] if item.get("candidate_index") is None] if offset == 0 else []
            body = {"source": record["source"], "candidates": batch, "media": assets, "source_issues": source_issues, "page_count": record["page_count"], "complete": offset + batch_size >= len(candidates)}
            self.request("POST", "/admin/mcq/import/batch", body)
        if candidates:
            final = self.request("POST", "/admin/mcq/import/finalize", {"file_hash": record["source"]["file_hash"], "source_path": record["source"]["source_path"]})
        else:
            final = self.request("POST", "/admin/mcq/import/finalize", {"file_hash": record["source"]["file_hash"], "source_path": record["source"]["source_path"], "processing_status": "REVIEW_REQUIRED" if record["issues"] else "NO_MCQS"})
        return final


def write_record(path: Path, record: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(path.suffix + ".tmp")
    with temp.open("w", encoding="utf-8", newline="") as stream:
        json.dump(record, stream, ensure_ascii=False)
        stream.write("\n")
    temp.replace(path)


def main() -> int:
    # Windows PowerShell may expose a cp1252 stdout even when the file paths
    # and question data contain Bangla. Never let a display encoding failure
    # abort the corpus audit after several thousand records have been written.
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("root", type=Path, help="Exam collection root directory")
    parser.add_argument("--output-dir", type=Path, help="Local JSONL and audit output directory")
    parser.add_argument("--api-url", default=os.getenv("SHIKHOK_API_URL"), help="Live app origin; omit to prepare an offline import pack")
    parser.add_argument("--email", default=os.getenv("SHIKHOK_ADMIN_EMAIL"))
    parser.add_argument("--password", default=os.getenv("SHIKHOK_ADMIN_PASSWORD"))
    parser.add_argument("--ocr-languages", default="ben+eng", help="Tesseract trained languages (default ben+eng)")
    parser.add_argument("--max-files", type=int, help="Process at most N files; useful for a pilot")
    parser.add_argument("--workers", type=int, default=min(4, max(1, os.cpu_count() or 1)), help="Maximum concurrent file conversions (default: up to 4)")
    parser.add_argument("--resume", action="store_true", help="Keep existing audit records and skip files with saved successful extraction records")
    parser.add_argument("--batch-size", type=int, default=20)
    args = parser.parse_args()
    root = args.root.resolve()
    if not root.is_dir():
        parser.error(f"source folder does not exist: {root}")
    output_dir = (args.output_dir or root.parent / f"{root.name}-mcq-import").resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    audit_path = output_dir / "audit.jsonl"
    previous_rows: dict[str, dict[str, Any]] = {}
    if args.resume and audit_path.exists():
        with audit_path.open("r", encoding="utf-8") as stream:
            for line in stream:
                if line.strip():
                    previous = json.loads(line)
                    previous_rows[previous["source_path"]] = previous
    else:
        audit_path.write_text("", encoding="utf-8")
    client = None
    if args.api_url:
        if not args.email or not args.password:
            parser.error("--api-url requires --email/--password or SHIKHOK_ADMIN_EMAIL/SHIKHOK_ADMIN_PASSWORD")
        client = ApiClient(args.api_url, args.email, args.password)
    files = [path for path in root.rglob("*") if path.is_file() and path.suffix.casefold() in SUPPORTED and output_dir not in path.parents]
    files.sort(key=lambda value: value.relative_to(root).as_posix().casefold())
    discovered_count = len(files)
    if args.max_files:
        files = files[: max(0, args.max_files)]
    if args.resume:
        saved_records: dict[str, dict[str, Any]] = {}
        for saved_path in (output_dir / "files").glob("*.json"):
            try:
                with saved_path.open("r", encoding="utf-8") as stream:
                    saved = json.load(stream)
                saved_records[saved.get("source", {}).get("source_path", "")] = saved
            except (OSError, json.JSONDecodeError):
                continue
        resumable_statuses = {"FAILED", "UPLOAD_FAILED"}
        skip_paths: set[str] = set()
        for path in files:
            relative = path.relative_to(root).as_posix()
            row = previous_rows.get(relative)
            record = saved_records.get(relative)
            if not row or row.get("processing_status") in resumable_statuses or not record:
                continue
            expected_hash = row.get("file_hash")
            saved_hash = record.get("source", {}).get("file_hash")
            if expected_hash and expected_hash == saved_hash == sha256_file(path):
                skip_paths.add(relative)
        files = [path for path in files if path.relative_to(root).as_posix() not in skip_paths]
        resumable_rows = [previous_rows[path] for path in skip_paths]
    else:
        saved_records = {}
        resumable_rows = []
    if not 1 <= args.workers <= 8:
        parser.error("--workers must be between 1 and 8")
    audits: list[dict[str, Any]] = sorted(resumable_rows, key=lambda row: row["source_path"].casefold())
    if args.resume:
        with audit_path.open("w", encoding="utf-8", newline="") as stream:
            for row in audits:
                stream.write(json.dumps(row, ensure_ascii=False) + "\n")
    tasks: Queue = Queue(maxsize=args.workers)
    results: Queue = Queue(maxsize=args.workers)
    workers = [Thread(target=_import_worker, args=(tasks, results, root, args.ocr_languages), name=f"mcq-import-{index + 1}", daemon=True) for index in range(args.workers)]
    for worker in workers:
        worker.start()
    pending_results: dict[int, tuple[dict[str, Any] | None, Exception | None, float]] = {}
    next_submit = 0
    for queued_index, queued_path in enumerate(files[: args.workers]):
        tasks.put((queued_index, queued_path))
        next_submit += 1
    for index, path in enumerate(files, start=1):
        task_index = index - 1
        while task_index not in pending_results:
            result_index, record_result, processing_error, extraction_seconds = results.get()
            pending_results[result_index] = (record_result, processing_error, extraction_seconds)
        record_result, processing_error, extraction_seconds = pending_results.pop(task_index)
        relative = path.relative_to(root).as_posix()
        audit: dict[str, Any] = {"source_path": relative, "processing_status": "FAILED", "error": None, "detected": 0, "local_valid_candidates": 0, "needs_review": 0, "source_issue_count": 0, "candidate_reconciliation_ok": False, "duplicates": 0, "imported": 0, "page_count": 0}
        try:
            if processing_error:
                raise processing_error
            if record_result is None:
                raise RuntimeError("file worker returned no extraction record")
            record = record_result
            review_candidates = {issue["candidate_index"] for issue in record["issues"] if issue.get("candidate_index") is not None}
            source_issue_count = sum(issue.get("candidate_index") is None for issue in record["issues"])
            detected = len(record["questions"]) + len(review_candidates)
            audit.update({"file_hash": record["source"]["file_hash"], "file_size": record["source"]["file_size"], "file_type": record["source"]["file_type"], "class_level": record["source"]["class_level"], "subject": record["source"]["subject"], "part": record["source"]["part"], "chapter": record["source"]["chapter"], "detected": detected, "local_valid_candidates": len(record["questions"]), "needs_review": len(review_candidates), "source_issue_count": source_issue_count, "candidate_reconciliation_ok": detected == len(record["questions"]) + len(review_candidates), "page_count": record["page_count"]})
            record_key = sha256_bytes(f"{relative}\0{record['source']['file_hash']}".encode("utf-8"))
            record_path = output_dir / "files" / f"{record_key}.json"
            write_record(record_path, record)
            if client:
                try:
                    result = client.send_file(record, max(1, args.batch_size))
                    audit.update({"imported": result.get("totalImportedMcqs", 0), "duplicates": result.get("totalDuplicates", 0), "needs_review": result.get("totalAmbiguous", audit["needs_review"]), "rejected": result.get("totalRejected", 0), "processing_status": result.get("processingStatus", "REVIEW_REQUIRED"), "server_reconciliation_ok": result.get("totalDetectedMcqs", 0) == sum(result.get(key, 0) for key in ("totalImportedMcqs", "totalDuplicates", "totalRejected", "totalAmbiguous"))})
                    audit["candidate_reconciliation_ok"] = audit["server_reconciliation_ok"]
                except Exception as error:
                    audit["error"] = str(error)
                    audit["processing_status"] = "UPLOAD_FAILED"
            else:
                audit["processing_status"] = "REVIEW_REQUIRED" if record["issues"] else "READY_TO_IMPORT"
        except Exception as error:
            audit["error"] = f"{type(error).__name__}: {error}"
            audit["processing_status"] = "FAILED"
            if client:
                source = {"source_path": relative, "source_filename": path.name, "file_type": path.suffix[1:].upper(), "file_size": path.stat().st_size, "file_hash": sha256_file(path), **infer_metadata(path, root)}
                try:
                    client.request("POST", "/admin/mcq/import/failure", {"source": source, "error": audit["error"]})
                except Exception as upload_error:
                    audit["error"] += f"; failure audit upload failed: {upload_error}"
        audit["extraction_seconds"] = round(extraction_seconds, 2)
        audits.append(audit)
        with audit_path.open("a", encoding="utf-8") as stream:
            stream.write(json.dumps(audit, ensure_ascii=False) + "\n")
        print(f"[{index}/{len(files)}] {audit['processing_status']}: {relative}", flush=True)
        if next_submit < len(files):
            queued_path = files[next_submit]
            tasks.put((next_submit, queued_path))
            next_submit += 1

    for _ in workers:
        tasks.put(None)
    for worker in workers:
        worker.join()
    totals = {key: sum(int(row.get(key, 0)) for row in audits) for key in ("detected", "imported", "duplicates", "rejected", "needs_review", "local_valid_candidates", "source_issue_count")}
    reconciliation = totals["detected"] == totals["imported"] + totals["duplicates"] + totals["rejected"] + totals["needs_review"] if client else None
    candidate_reconciliation = (
        not any(row["processing_status"] in {"FAILED", "UPLOAD_FAILED"} for row in audits)
        and (reconciliation if client else totals["detected"] == totals["local_valid_candidates"] + totals["needs_review"])
    )
    summary = {"created_at": datetime.now(timezone.utc).isoformat(), "root": str(root), "output_dir": str(output_dir), "files_discovered": discovered_count, "files_processed": sum(row["processing_status"] not in {"FAILED", "UPLOAD_FAILED"} for row in audits), "files_failed": sum(row["processing_status"] in {"FAILED", "UPLOAD_FAILED"} for row in audits), "totals": totals, "candidate_reconciliation_ok": candidate_reconciliation, "server_reconciliation_ok": reconciliation, "unresolved_files": [row["source_path"] for row in audits if row["processing_status"] not in {"COMPLETED", "NO_MCQS"}]}
    write_record(output_dir / "summary.json", summary)
    with (output_dir / "audit.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(audits[0].keys()) if audits else ["source_path", "processing_status"])
        writer.writeheader()
        writer.writerows(audits)
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    close_word()
    return 0 if not summary["files_failed"] and (reconciliation is not False) else 2


if __name__ == "__main__":
    sys.exit(main())
